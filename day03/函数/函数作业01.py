def char_count(text:str)->dict:
    count = {}
    for char in text:
        if count.get(char) is None:
            count[char] = 1
        else:
            count[char] += 1
    return count

count = char_count("asdfasfsadfdasfdasx")
print(count)

print("====================================================")

import os
def select_content(file_path,key):
    if not os.path.exists(file_path):
        return
    data_list = []
    with open(file_path,'r',encoding='utf-8') as f:
        for line in f:
            if key in line:
                data_list.append(line.strip("\n"))
    return data_list

result = select_content('./aa.txt','股票')
if result is None:
    print("文件不存在")
else:
    print(result)

print("---------------------")

import hashlib
from openpyxl import load_workbook

def get_user_dict():
    user_dict = {}
    wb = load_workbook('./user.xlsx')
    sheet = wb.worksheets[0]
    for row in sheet.iter_rows(min_row=2):
        user_dict[row[1].value] = row[2].value
    return user_dict

def encrypt(origin):
    origin_bytes = origin.encode('utf-8')
    md5_object = hashlib.md5()
    md5_object.update(origin_bytes)
    return md5_object.hexdigest()

user = input(">>>用户名：").strip()
pwd = input(">>>密码:").strip()
encrypt_pwd = encrypt(pwd)
user_dict = get_user_dict()
ob_pwd = user_dict.get(user)
if encrypt_pwd == ob_pwd:
    print("登陆成功")
else:
    print("登陆失败")

print("---------------------")

