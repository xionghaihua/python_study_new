import requests
from openpyxl import load_workbook
import os
SAVE_DIR="downloads"
def download(url):
    res = requests.get(
        url=url,
        headers={
            "user-agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/147.0.0.0 Safari/537.36"
        }
    )
    title = url.split("/")[-1].split(".")[0]
    file_type = url.split("/")[-1].split(".")[1]
    filename='{}.{}'.format(title,file_type)
    filepath = os.path.join(SAVE_DIR, filename)
    with open(filepath, 'wb') as f:
        f.write(res.content)

    print(f"已下载: {filename}")

def main():
    os.makedirs(SAVE_DIR, exist_ok=True)
    wb = load_workbook("工作簿.xlsx")
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        print(f"--- 正在处理工作表: {sheet_name} ---")
        for row in range(3,sheet.max_row+1):
            url = sheet.cell(row=row, column=1).value
            if url and str(url).startswith("http"):
                download(str(url))
if __name__ == "__main__":
    main()

