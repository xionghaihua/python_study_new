from openpyxl import Workbook
from openpyxl.styles import PatternFill,Font,Alignment,Border,Side
wb = Workbook()
ws = wb.active

title_font = Font(name='微软雅黑', size=16, bold=True, color='FFFFFF')
title_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
title_align = Alignment(horizontal='center', vertical='center')

header_font = Font(name='宋体', size=12, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='5B9BD5', end_color='5B9BD5', fill_type='solid')
header_border = Border(
      left=Side(style='thin'),
      right=Side(style='thin'),
      top=Side(style='thin'),
      bottom=Side(style='thin')
)
# 数据样式
data_border = Border(
      left=Side(style='thin'),
      right=Side(style='thin'),
      top=Side(style='thin'),
      bottom=Side(style='thin')
)
#单元格合并
ws.merge_cells("A1:E1")
#解除合并 ws.unmerge_cells("A1:E1")
ws['A1'].value = '销售数据表'
ws['A1'].font = title_font
ws['A1'].fill = title_fill
ws['A1'].alignment = title_align

# 设置表头
headers = ['姓名', '部门', '销售额', '完成率', '状态']
for col,header in enumerate(headers,start=1):
    cell = ws.cell(row=2,column=col,value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.border = header_border
    cell.alignment = Alignment(horizontal='center', vertical='center')

# 设置数据
data = [
      ['张三', '销售一部', 120000, '120%', '优秀'],
      ['李四', '销售二部', 95000, '95%', '良好'],
      ['王五', '销售一部', 80000, '80%', '达标'],
]
for row_idx,row_data in enumerate(data,start=3):
    for col_idx,value in enumerate(row_data,start=1):
        cell = ws.cell(row=row_idx,column=col_idx,value=value)
        cell.border = data_border
        cell.alignment = Alignment(horizontal='center', vertical='center')
        # 状态列不同颜色
        if col_idx == 5:
            if value == '优秀':
                cell.fill = PatternFill(start_color='00B050', end_color='00B050',fill_type='solid')
            elif value == '良好':
                cell.fill = PatternFill(start_color='92D050', end_color='92D050',
  fill_type='solid')
            else:
                cell.fill = PatternFill(start_color='FFFF00', end_color='FFFF00',
  fill_type='solid')

#设置单元格的宽高
ws.row_dimensions[1].height = 30
ws.column_dimensions['A'].width = 20


#公式


wb.save('./p3.xlsx')


