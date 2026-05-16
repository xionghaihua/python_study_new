#在原Excel文件基础上写内容
"""
from openpyxl import load_workbook
wb = load_workbook("p1.xlsx")
ws = wb["数据导出"]
cell = ws.cell(row=1,column=4)
cell.value = "姓名"
wb.save("p1.xlsx")
"""

#新创建excel
"""
from openpyxl import workbook
wb = workbook.Workbook()
"""
#操作sheet操作
"""
from openpyxl import workbook
wb = workbook.Workbook()

#修改sheet名称
ws = wb.worksheets[0]
ws.title = "数据集"
wb.save("./p2.xlsx")

#创建sheet并设置sheet颜色
sheet = wb.create_sheet("工作计划",0) #0,表示放到最开始的位置
sheet.sheet_properties.tabColor = "1072BA"
wb.save('./p2.xlsx')
"""

#修改表格数据

from openpyxl import load_workbook
wb2 = load_workbook('./p1.xlsx')
ws2 = wb2["数据导出"]

"""
#获取单个单元格，修改值
cell = ws2.cell(row=2,column=5)
cell.value = "否"
wb2.save('./p1.xlsx')

#获取某些单元格
cell_list = ws2["B2":"C3"]
#print(cell_list) #((<Cell '数据导出'.B2>, <Cell '数据导出'.C2>), (<Cell '数据导出'.B3>, <Cell '数据导出'.C3>))
for row in cell_list:
    for cell in row:
        cell.value = "新值"
wb2.save('./p1.xlsx')
"""


#单元格样式操作
"""
#对齐方式
from openpyxl.styles import Alignment, Border,Side,Font,PatternFill,GradientFill
cell = ws2.cell(row=1,column=4)
#horizontal 水平
#vertical 垂直
#text_rotation 旋转角度
#wrap_text是否自动换行
cell.alignment = Alignment(horizontal="center", vertical="center",text_rotation=0,wrap_text=True)
wb2.save('./p1.xlsx')

#边框
#边框样式: dashDot,dashDotDot,dashed ,dotted,double,hair,medium,mediumDashDot,mediumDashDotDot,mediumDashed,thick,thin

cell = ws2.cell(row=1,column=4)
print(cell.value)
cell.border = Border(
    top=Side(style='dashDot',color="FFB6C1"),
    bottom=Side(style='dashDot',color="9932CC"),
    left=Side(style='dashDot',color="9932CC"),
    right=Side(style='dashDot',color="9932CC"),
)
wb2.save('./p1.xlsx')
"""

from openpyxl import load_workbook
from openpyxl.styles import PatternFill,GradientFill,Font,Border,Side,Alignment
#字体样式
wb3 = load_workbook('./p1.xlsx')
ws3 = wb3["数据导出"]
#
ws3['D1'].font = Font(
    name='微软雅黑',
    size=16,
    bold=True,
    italic=False, #斜体
    underline='single', #下划线 single，double
    strike=False,  #删除线
    color='FF0000',  #颜色
)

#对齐方式
ws3['C2'].alignment = Alignment(
    horizontal="center",
    vertical="center",
    wrap_text=True,       # 自动换行
    shrink_to_fit=False,  # 缩小字体填充
    indent=1,             # 缩进
    text_rotation=0     # 文字旋转角度
)


#表格边框
thin_border = Border(
    left=Side(style='thin', color='000000'),
    right=Side(style='thin', color='000000'),
    top=Side(style='thin', color='000000'),
    bottom=Side(style='thin', color='000000')
)
for row in range(1,6):
    for col in range(1,6):
        ws3.cell(row=row,column=col).border = thin_border


#背景填充
pattern = PatternFill(
    start_color='4472C4',
    end_color='4472C4',
    fill_type='solid'
)
for row in range(1,6):
    for col in range(1,6):
        ws3.cell(row=row,column=col).fill = pattern
wb3.save('./p1.xlsx')


